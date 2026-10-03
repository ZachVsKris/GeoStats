"""Extract traceable candidates without averaging, imputing or gameplay approval."""
import collections,csv,gzip,hashlib,io,json,math,re,zipfile
from pathlib import Path
import openpyxl,xlrd
ROOT=Path(__file__).resolve().parents[1];S=ROOT/'data/animalstats/source';R=ROOT/'data/animalstats/research'
observations=[]
def numeric(x):
 try:
  n=float(x)
  return n if math.isfinite(n) and n>=0 else None
 except (ValueError,TypeError):return None
def add(concept,species,value,unit,file,row,reference='',conditions=None,gate='Primary measurement, uncertainty and comparability audit required.'):
 n=numeric(value)
 if n is None or not species:return
 observations.append(dict(concept=concept,species=str(species).strip().replace('_',' '),value=n,unit=unit,file=file,row=row,reference=str(reference),conditions=conditions or {},gate=gate,review_status='pending',game_approved=False))
for key,concept,unit in [('fh','heart-rate','beats/min'),('fr','breathing-rate','breaths/min'),('sv','heart-stroke-volume','mL/beat'),('vt','breath-volume','mL/breath')]:
 p=S/f'discovery-heart-{key}data-filter.csv'
 for i,r in enumerate(csv.DictReader(p.open()),2):
  if not all(r.get(k)==v for k,v in [('adult','y'),('resting','y'),('sedated','n')]):continue
  add(concept,' '.join([r.get('genus',''),r.get('species','')]),r.get(key),unit,p.name,i,r.get('Primary Source',r.get('','')),r)
p=S/'discovery-field-energy-FieldMR_SciDataPub.xls';w=xlrd.open_workbook(p);sh=w.sheet_by_name('Alldata');headers=sh.row_values(0)
for i in range(1,sh.nrows):
 r=dict(zip(headers,sh.row_values(i)))
 if r['Kingdom']!='Animalia':continue
 add('field-energy',r['SpeciesAcceptedName'],r['FMR_kJ_d'],'kJ/day',p.name,i+1,r['Reference'],r,'BLOCKED: source methods include range midpoints and estimated mass conversions; original-reference audit required. Routine laboratory oxygen measurements and field measurements must be separated.')
p=S/'discovery-primate-vocabulary-ZenodoData_April2026.xlsx';w=openpyxl.load_workbook(p,read_only=True,data_only=True);sh=list(w['data'].values);h=sh[0]
for i,line in enumerate(sh[1:],2):
 r=dict(zip(h,line));add('vocal-repertoire',r.get('BinomialNomenclature'),r.get('Updatedrepertoire size'),'adult call types',p.name,i,r.get('RepertoireUpdateReference'),{'family':r.get('Family')},'Adult repertoire definitions and observation effort require audit; no ancestor/model values extracted.')
p=next(S.glob('discovery-whiskers-*.xlsx'));w=openpyxl.load_workbook(p,read_only=True,data_only=True)
for sheet in ['Flat Floor','Inclined Floor']:
 rows=list(w[sheet].values);h=rows[1]
 for i,line in enumerate(rows[2:],3):
  r=dict(zip(h,line));add('whisker-length',r.get('Species'),r.get('Whisker Length'),'UNVERIFIED',p.name,i,'https://zenodo.org/records/4942927',{'sheet':sheet,'clip':r.get('Clip')},'BLOCKED: unit and whether repeated clips reuse specimen measurements require verification; do not aggregate clips as independent animals.')
p=S/'discovery-fish-reproduction.tsv';w=openpyxl.load_workbook(io.BytesIO(p.read_bytes()),read_only=True,data_only=True)
for i,r in enumerate(list(w['Traits value'].values)[1:],2):
 if r[1]!='S' or re.search('infer|estimate|assum',str(r[12]),re.I):continue
 add('offspring-count',r[0],r[8],'reported fecundity',p.name,i,r[9],{'habitat':r[4],'care':r[5],'remark':r[12]},'BLOCKED: metadata mixes eggs per year and one ovarian batch; cannot treat as babies at once. Original references required.')
 add('egg-length',r[0],r[10],'mm',p.name,i,r[11],{'care':r[5]},'BLOCKED: offspring size mixes eggs and live young; identify anatomy before ranking.')
p=next(S.glob('discovery-mammal-parenting-*.xlsx'));w=openpyxl.load_workbook(p,read_only=True,data_only=True);rows=list(w.active.values);h=rows[0]
for i,line in enumerate(rows[1:],2):
 r=dict(zip(h,line))
 for field,concept,unit in [('LitterSize','offspring-count','offspring/litter'),('NeonateMass','newborn-mass','UNVERIFIED'),('LittersYear','reproduction-frequency','litters/year')]:
  add(concept,r.get('Species'),r.get(field),unit,p.name,i,'https://zenodo.org/records/4935027',{},'Compiled averages; verify original source values, sex/stage and units. AnnualFecundity and TeatRatio excluded as derived endpoints.')
p=S/'discovery-arthropods.zip';z=zipfile.ZipFile(p)
taxa={r['id']:r for r in csv.DictReader(io.StringIO(z.read('taxon.txt').decode()),delimiter='\t')}
for i,r in enumerate(csv.DictReader(io.StringIO(z.read('measurementorfacts.txt').decode()),delimiter='\t'),2):
 t=taxa.get(r['id'],{})
 mapping={'Body_size':'length','Fecundity':'offspring-count','Lifespan':'lifespan','Development_time':'maturity-age'}
 if r['measurementType'] not in mapping:continue
 add(mapping[r['measurementType']],t.get('scientificName'),r['measurementValue'],r['measurementUnit'],p.name,i,'https://doi.org/10.3897/BDJ.13.e146785',t,'BLOCKED: literature and expert knowledge mixed without per-cell method/source in archive; CC-BY-NC license. Body-size anatomical definition, fecundity time basis and development versus maturity need verification. Thermal niche and dispersal score excluded.')
# Structured account facts: accept explicit scalars only; ranges stay in the
# original account snapshot and are not silently replaced by midpoints.
mapping={'mass':'mass','length':'length','gestation period':'pregnancy-duration','weaning age':'milk-duration','number of offspring':'offspring-count','lifespan':'lifespan','sexual or reproductive maturity':'maturity-age','basal metabolic rate':'resting-energy'}
for fname in ['adw-measurement-accounts.json','expansion-adw-accounts.json']:
 p=S/fname
 if not p.exists():continue
 for a in json.loads(p.read_text()):
  for section in a.get('sections',{}).values():
   for label,text in section.get('items',{}).items():
    concept=next((v for k,v in mapping.items() if k in label.lower()),None)
    m=re.fullmatch(r'([0-9]+(?:\.[0-9]+)?)\s*([^0-9]*)',text.strip())
    if not concept or not m:continue
    add(concept,a['url'].rstrip('/').split('/')[-1],m[1],m[2].strip() or 'UNVERIFIED',p.name,label,a['url'],{'reported_label':label},'BLOCKED: account-level summary; trace primary citation and preserve average versus record, wild versus captive, sex and taxonomy. No prose copied into ranking data.')
with gzip.open(R/'expansion-evidence-records.jsonl.gz','wt') as f:
 for o in observations:f.write(json.dumps(o)+'\n')
summary=[]
for concept in sorted({o['concept'] for o in observations}):
 rows=[o for o in observations if o['concept']==concept]
 summary.append(dict(concept=concept,numeric_rows=len(rows),source_taxon_labels=len({o['species'] for o in rows}),units='; '.join(sorted({o['unit'] for o in rows})),files='; '.join(sorted({o['file'] for o in rows})),primary_reference_rows=sum(bool(o['reference']) and o['reference']!='None' for o in rows),certified_board_ready=False,remaining_gates='; '.join(sorted({o['gate'] for o in rows}))))
with (R/'expansion-endpoint-audit.csv').open('w',newline='') as f:
 writer=csv.DictWriter(f,fieldnames=list(summary[0]));writer.writeheader();writer.writerows(summary)
report=dict(extracted_numeric_rows=len(observations),concepts_with_extracted_values=len(summary),certified_board_ready=0,new_playable_approvals=0,source_files={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in S.iterdir() if p.name in {o['file'] for o in observations}},endpoints=summary)
(R/'expansion-evidence-summary.json').write_text(json.dumps(report,indent=2)+'\n');print(json.dumps({k:v for k,v in report.items() if k not in ('source_files','endpoints')},indent=2))
