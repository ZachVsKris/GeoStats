#!/usr/bin/env python3
"""Pin AnthropInsect's original release in a separate, unapproved research store.

This is deliberately not an importer for pilot.json. Reported extrema still need
their original species-level study, stage, caste and measurement basis reviewed.
"""
import argparse
import gzip
import hashlib
import json
import re
import sqlite3
from collections import Counter
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
FAMILIES = {
    'Body_size_adult': ('BS', 'mm', 'Longest adult body', 'Shortest adult body'),
    'Lifespan': ('LS', 'days', 'Longest life', 'Shortest life'),
    'Fecundity': ('Fec', 'eggs/female/lifetime', 'Most eggs in a lifetime', 'Fewest eggs in a lifetime'),
    'Voltinism': ('Vol', 'generations/year', 'Most generations per year', 'Fewest generations per year'),
}
COMMON = {
    'Apis mellifera', 'Danaus plexippus', 'Bombyx mori', 'Gryllus bimaculatus',
    'Locusta migratoria', 'Schistocerca gregaria', 'Blattella germanica',
    'Periplaneta americana', 'Coccinella septempunctata', 'Tenebrio molitor',
    'Acheta domesticus', 'Drosophila melanogaster', 'Bombus terrestris',
    'Mantis religiosa', 'Lasius niger', 'Musca domestica',
}

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--workbook', required=True, type=Path)
    parser.add_argument('--metadata', required=True, type=Path)
    parser.add_argument('--output', type=Path, default=ROOT / 'data/animalstats/research/anthropinsect')
    args = parser.parse_args()
    metadata = json.loads(args.metadata.read_text())
    version = metadata['data']['latestVersion'] if 'data' in metadata else {
        **metadata, 'files': [{'dataFile': metadata['dataFile']}],
    }
    file = next(f['dataFile'] for f in version['files'] if f['dataFile']['filename'] == args.workbook.name)
    content = args.workbook.read_bytes()
    assert len(content) == file['filesize'], 'Original workbook size mismatch'
    assert hashlib.md5(content).hexdigest() == file['md5'], 'Original workbook checksum mismatch'
    assert version['license']['name'] == 'CC BY 4.0'
    book = openpyxl.load_workbook(args.workbook, read_only=True, data_only=True)
    rows = book['DB'].iter_rows(values_only=True)
    headers = next(rows)
    species = [(i, dict(zip(headers, r))) for i, r in enumerate(rows, 2) if r[9] and str(r[9]).strip()]
    definitions = {str(r[0]).strip(): {'definition': r[1], 'units': r[2], 'abbreviation': r[3]}
                   for r in book['Variables'].iter_rows(values_only=True) if r[0]}
    refs = []
    for i, r in enumerate(book['References'].iter_rows(values_only=True), 1):
        if isinstance(r[1], (int, float)) and r[2] and r[4]:
            refs.append({'row': i, 'id': r[1], 'species': str(r[2]), 'traits': str(r[3] or ''),
                         'citation': str(r[4]), 'sisterSpecies': str(r[5] or ''), 'details': r[6]})
    # A reference mentioning a species is a lead, not proof that a particular
    # cell comes from that study. The workbook does not map every cell to a ref.
    exact = {}
    for ref in refs:
        for taxon in re.split(r'[/\n;]', ref['species']):
            exact.setdefault(taxon.strip(), []).append(ref)
    args.output.mkdir(parents=True, exist_ok=True)
    database = args.output / 'research.sqlite'
    temporary = args.output / 'research.sqlite.tmp'
    temporary.unlink(missing_ok=True)
    connection = sqlite3.connect(temporary)
    connection.executescript('''
        CREATE TABLE metadata(key TEXT PRIMARY KEY, value TEXT NOT NULL);
        CREATE TABLE species(source_row INTEGER PRIMARY KEY, scientific_name TEXT NOT NULL, raw_json TEXT NOT NULL);
        CREATE TABLE references_raw(source_row INTEGER PRIMARY KEY, raw_json TEXT NOT NULL);
        CREATE TABLE observations(source_row INTEGER NOT NULL, scientific_name TEXT NOT NULL,
          family TEXT NOT NULL, source_field TEXT NOT NULL, raw_value TEXT NOT NULL,
          unit TEXT NOT NULL, status TEXT NOT NULL CHECK(status != 'approved'),
          review_reasons TEXT NOT NULL, reference_leads TEXT NOT NULL,
          PRIMARY KEY(source_row, source_field));
        CREATE INDEX observations_taxon ON observations(scientific_name, family);
    ''')
    counts, coverage, anomalies = Counter(), [], []
    for row, record in species:
        name = str(record['Species']).strip()
        connection.execute('INSERT INTO species VALUES (?,?,?)', (row, name, json.dumps(record, ensure_ascii=False)))
        available = []
        for family, (abbr, unit, high, low) in FAMILIES.items():
            leads = [r for r in exact.get(name, []) if any(token in {'bs', 'bsa'} if abbr == 'BS' else token == abbr.lower() for token in re.split(r'[/,;\s]+', r['traits'].lower()))]
            for suffix in ['min', 'max', 'estimate', 'mean']:
                field = family + '_' + suffix
                value = record.get(field)
                if value is None or value == '':
                    continue
                reasons = ['Original study, species, stage/caste, sex and conditions must be verified.',
                           'Reference mentions are leads; no per-cell citation mapping is supplied.']
                status = 'needs-primary-review'
                if suffix == 'estimate':
                    status = 'excluded-estimate'
                    reasons.append('Typical estimate is not a measured value.')
                elif suffix == 'mean':
                    status = 'excluded-derived-mean'
                    reasons.append('Workbook defines mean as estimate and/or midpoint of min and max.')
                if not leads:
                    reasons.append('No exact-species reference lead matched this trait abbreviation.')
                if any(r['sisterSpecies'].strip().lower() != 'no' for r in leads):
                    reasons.append('At least one reference lead has missing/positive sister-species estimation flag.')
                counts[status] += 1
                counts[family + ':' + suffix] += 1
                connection.execute('INSERT INTO observations VALUES (?,?,?,?,?,?,?,?,?)',
                    (row, name, family, field, str(value), unit, status, json.dumps(reasons), json.dumps(leads)))
                available.append(field)
            lo, hi, estimate = (record.get(family + '_' + s) for s in ['min', 'max', 'estimate'])
            if isinstance(lo, (int, float)) and isinstance(hi, (int, float)) and lo > hi:
                anomalies.append({'sourceRow': row, 'species': name, 'family': family, 'issue': 'reversed-range', 'min': lo, 'max': hi})
            if all(isinstance(v, (int, float)) for v in [lo, hi, estimate]) and not lo <= estimate <= hi:
                anomalies.append({'sourceRow': row, 'species': name, 'family': family,
                                  'issue': 'estimate-outside-reported-range', 'min': lo, 'max': hi, 'estimate': estimate})
        if name in COMMON:
            coverage.append({'sourceRow': row, 'species': name, 'fields': available,
                             'referenceLeads': exact.get(name, []), 'status': 'not-approved'})
    for ref in refs:
        connection.execute('INSERT INTO references_raw VALUES (?,?)', (ref['row'], json.dumps(ref)))
    manifest = {
        'dataset': 'AnthropInsect 2.0', 'doi': '10.48579/PRO/LYB70N',
        'sourceUrl': 'https://data.indores.fr/api/access/datafile/37049',
        'releaseTime': version['releaseTime'], 'license': version['license']['name'],
        'sha256': hashlib.sha256(content).hexdigest(), 'md5': file['md5'],
        'speciesRows': len(species), 'uniqueSpecies': len({r['Species'] for i, r in species}),
        'referenceRows': len(refs), 'observations': sum(v for k, v in counts.items() if ':' not in k),
        'counts': dict(counts), 'approvedObservations': 0,
        'gameplayLabelsAdded': 0,
        'labelPairsToReview': [{'family': f, 'labels': [s[2], s[3]]} for f, s in FAMILIES.items()],
        'admissionRules': [
            'Never approve typical estimates, calculated means, genus/family substitutions or sister-species values.',
            'Verify reported min/max against the exact primary species study before admission.',
            'Queen, worker and other caste measurements cannot be silently mixed.',
            'Eggs per day or clutch cannot be presented as eggs over a lifetime.',
            'Body length must not use caterpillar length, wingspan, antennae or appendages.',
            'Generations per year must state geography and conditions; not a universal species constant.',
            'Latitude/longitude/climate extents of occurrence records are not complete biological range maps.',
        ],
        'anomalyCount': len(anomalies), 'commonSpeciesFound': len(coverage),
    }
    for k, v in manifest.items():
        connection.execute('INSERT INTO metadata VALUES (?,?)', (k, json.dumps(v)))
    connection.commit()
    assert connection.execute('PRAGMA integrity_check').fetchone()[0] == 'ok'
    assert connection.execute("SELECT count(*) FROM observations WHERE status='approved'").fetchone()[0] == 0
    connection.close()
    temporary.replace(database)
    # Keep a compact, reproducible checkpoint in git; the working SQLite file
    # is regenerated and excluded from deployment/source control.
    (args.output / 'research.sqlite.gz').write_bytes(gzip.compress(database.read_bytes(), mtime=0))
    for name, obj in [('manifest.json', manifest), ('common-species-review.json', coverage),
                      ('range-anomalies.json', anomalies), ('variable-definitions.json', definitions)]:
        (args.output / name).write_text(json.dumps(obj, indent=2, ensure_ascii=False) + '\n')
    print(json.dumps(manifest, indent=2))

if __name__ == '__main__':
    main()
